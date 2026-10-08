"""Turn uploaded files into knowledge chunks.

``extract_file(data, filename, ext)`` returns an ``Extraction`` containing a list
of chunks (``{"text", "metadata", "key"}``) plus counters/notes for the source.
"""
import csv
import io
import json
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional

from app.chunkers.csv import chunk_csv
from app.chunkers.json import chunk_json
from app.chunkers.text import chunk_text, normalize_text
from app.training.records import (
    chunk_lines, find_records, looks_like_tickets, normalize_records,
)

TEXT_CHUNK_SIZE = 1000
TEXT_CHUNK_OVERLAP = 150

SUPPORTED_EXTENSIONS = {"pdf", "xml", "txt", "md", "csv", "json", "docx"}


def docx_available() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except Exception:
        return False


class ExtractionError(Exception):
    """The file could not be read (corrupt, encrypted, no text...)."""


@dataclass
class Extraction:
    chunks: List[Dict[str, Any]] = field(default_factory=list)
    records_count: int = 0          # tickets/records found (structured files)
    skipped_records: int = 0
    document_source_type: Optional[str] = None  # overrides documents.source_type (e.g. support_tickets)
    notes: List[str] = field(default_factory=list)


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    try:
        return data.decode("cp1252")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")


def _chunk(text: str, label: str, metadata: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Chunk prose and prefix each chunk with a short source label."""
    out = []
    for piece in chunk_text(text, TEXT_CHUNK_SIZE, TEXT_CHUNK_OVERLAP):
        out.append({"text": f"[{label}]\n{piece}" if label else piece, "metadata": dict(metadata), "key": None})
    return out


def _items_to_chunks(items: List[Dict[str, Any]], base_meta: Dict[str, Any]) -> List[Dict[str, Any]]:
    chunks = []
    for item in items:
        n = len(item["chunks"])
        for i, text in enumerate(item["chunks"]):
            meta = {**base_meta, **{k: v for k, v in item["metadata"].items() if v is not None}}
            if n > 1:
                meta["part"] = i + 1
            chunks.append({"text": text, "metadata": meta, "key": item["key"]})
    return chunks


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def extract_pdf(data: bytes, name: str) -> Extraction:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise ExtractionError("PDF support is not installed on the server (pypdf)") from e
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as e:
                raise ExtractionError("The PDF is password-protected") from e
        pages = reader.pages
        total = len(pages)
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Could not read PDF: {e}") from e

    ex = Extraction()
    empty_pages = 0
    for idx in range(total):
        try:
            text = pages[idx].extract_text() or ""
        except Exception:
            text = ""
        if not normalize_text(text):
            empty_pages += 1
            continue
        ex.chunks.extend(_chunk(text, f"{name} - Page {idx + 1}", {"filename": name, "page": idx + 1}))
    if not ex.chunks:
        raise ExtractionError(
            "No extractable text found in the PDF (it may be a scanned/image-only document; OCR is not supported)"
        )
    if empty_pages:
        ex.notes.append(f"{empty_pages} of {total} pages had no extractable text and were skipped")
    return ex


# ---------------------------------------------------------------------------
# XML
# ---------------------------------------------------------------------------

def _parse_xml(data: bytes):
    try:
        from defusedxml.ElementTree import fromstring  # safe against entity expansion attacks
    except ImportError:  # pragma: no cover
        from xml.etree.ElementTree import fromstring
    try:
        return fromstring(data)
    except Exception as e:
        raise ExtractionError(f"Invalid XML: {e}") from e


def _tag(el) -> str:
    tag = el.tag if isinstance(el.tag, str) else "node"
    return tag.split("}", 1)[-1]


def _xml_lines(el, path: str, out: List[str]) -> None:
    for attr, value in el.attrib.items():
        out.append(f"{path}@{attr.split('}', 1)[-1]}: {value}")
    text = (el.text or "").strip()
    if text:
        out.append(f"{path}: {' '.join(text.split())}")
    for child in el:
        _xml_lines(child, f"{path}/{_tag(child)}", out)
        tail = (child.tail or "").strip()
        if tail:
            out.append(f"{path}: {' '.join(tail.split())}")


def _xml_to_dict(el) -> Any:
    children = list(el)
    if not children and not el.attrib:
        return (el.text or "").strip()
    d: Dict[str, Any] = {k.split("}", 1)[-1]: v for k, v in el.attrib.items()}
    for child in children:
        key = _tag(child)
        val = _xml_to_dict(child)
        if key in d:
            if not isinstance(d[key], list):
                d[key] = [d[key]]
            d[key].append(val)
        else:
            d[key] = val
    text = (el.text or "").strip()
    if text:
        d["text"] = text
    return d


def _find_record_group(el, depth: int = 0):
    """Find the element whose repeated children are the records. Returns (parent, record_tag) or None."""
    children = list(el)
    if not children or depth > 6:
        return None
    counts: Dict[str, int] = {}
    for c in children:
        if len(c) or c.attrib:  # structured children only
            counts[_tag(c)] = counts.get(_tag(c), 0) + 1
    if counts:
        tag, n = max(counts.items(), key=lambda kv: kv[1])
        if n >= 2:
            return el, tag
    for c in children:
        if len(c):
            found = _find_record_group(c, depth + 1)
            if found:
                return found
    return None


def extract_xml(data: bytes, name: str) -> Extraction:
    root = _parse_xml(data)
    ex = Extraction()
    group = _find_record_group(root)
    if group is None:
        lines: List[str] = []
        _xml_lines(root, _tag(root), lines)
        if not lines:
            raise ExtractionError("The XML file contains no text")
        for c in chunk_lines(lines, header=f"[{name}]", max_chars=TEXT_CHUNK_SIZE + 200):
            ex.chunks.append({"text": c, "metadata": {"filename": name}, "key": None})
        return ex

    parent, rec_tag = group
    records = [c for c in parent if _tag(c) == rec_tag]
    ex.records_count = len(records)

    as_dicts = [_xml_to_dict(r) for r in records]
    if looks_like_tickets(as_dicts):
        items, skipped = normalize_records(as_dicts, "support_tickets")
        ex.skipped_records = skipped
        ex.document_source_type = "support_tickets"
        ex.chunks.extend(_items_to_chunks(items, {"filename": name}))
    else:
        for i, rec in enumerate(records):
            lines: List[str] = []
            _xml_lines(rec, rec_tag, lines)
            if not lines:
                ex.skipped_records += 1
                continue
            header = f"[{name} - {rec_tag} {i + 1}]"
            for c in chunk_lines(lines, header=header, max_chars=TEXT_CHUNK_SIZE + 200):
                ex.chunks.append({"text": c, "metadata": {"filename": name, "record": i + 1}, "key": None})

    # Non-record content (document headers etc.) kept as its own chunk(s)
    other: List[str] = []
    root_path = _tag(root)
    for attr, value in root.attrib.items():
        other.append(f"{root_path}@{attr}: {value}")
    for child in root:
        if child is parent or (parent is root and _tag(child) == rec_tag):
            continue
        if parent is not root and _contains(child, parent):
            continue
        _xml_lines(child, f"{root_path}/{_tag(child)}", other)
    if other:
        for c in chunk_lines(other, header=f"[{name}]", max_chars=TEXT_CHUNK_SIZE + 200):
            ex.chunks.append({"text": c, "metadata": {"filename": name}, "key": None})
    if not ex.chunks:
        raise ExtractionError("The XML file contains no text")
    return ex


def _contains(el, target) -> bool:
    return any(e is target for e in el.iter())


# ---------------------------------------------------------------------------
# DOCX / text / CSV / JSON
# ---------------------------------------------------------------------------

def extract_docx(data: bytes, name: str) -> Extraction:
    try:
        import docx
    except ImportError as e:
        raise ExtractionError("DOCX support is not installed on the server (python-docx)") from e
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractionError(f"Could not read DOCX: {e}") from e
    parts: List[str] = []
    for p in document.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append("\n".join(rows))
    text = "\n\n".join(parts)
    if not text.strip():
        raise ExtractionError("The DOCX file contains no text")
    ex = Extraction()
    ex.chunks = _chunk(text, name, {"filename": name})
    return ex


def extract_plain(data: bytes, name: str) -> Extraction:
    text = decode_text(data)
    if not text.strip():
        raise ExtractionError("The file is empty")
    ex = Extraction()
    ex.chunks = _chunk(text, name, {"filename": name})
    return ex


def _ticket_extraction(records: List[Any], name: str) -> Extraction:
    ex = Extraction(records_count=len(records), document_source_type="support_tickets")
    items, skipped = normalize_records(records, "support_tickets")
    ex.skipped_records = skipped
    ex.chunks = _items_to_chunks(items, {"filename": name})
    return ex


def extract_csv(data: bytes, name: str) -> Extraction:
    text = decode_text(data)
    if not text.strip():
        raise ExtractionError("The file is empty")
    try:
        rows = list(csv.DictReader(io.StringIO(text)))
    except Exception:
        rows = []
    if rows and looks_like_tickets(rows):
        return _ticket_extraction(rows, name)

    ex = Extraction(records_count=max(0, len(rows)))
    header = text.splitlines()[0] if text.splitlines() else ""
    for i, chunk in enumerate(chunk_csv(text, rows_per_chunk=20)):
        body = chunk if i == 0 or not header else f"{header}\n{chunk}"
        for piece in chunk_lines(body.splitlines(), header=f"[{name}]", max_chars=2000):
            ex.chunks.append({"text": piece, "metadata": {"filename": name}, "key": None})
    return ex


def extract_json(data: bytes, name: str) -> Extraction:
    text = decode_text(data)
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as e:
        raise ExtractionError(f"Invalid JSON: {e}") from e
    try:
        records, _ = find_records(parsed)
    except ValueError:
        records = []
    if records and looks_like_tickets(records):
        return _ticket_extraction(records, name)
    ex = Extraction(records_count=len(records) if isinstance(parsed, (list, dict)) else 0)
    for chunk in chunk_json(text):
        ex.chunks.append({"text": f"[{name}]\n{chunk}", "metadata": {"filename": name}, "key": None})
    return ex


EXTRACTORS = {
    "pdf": extract_pdf,
    "xml": extract_xml,
    "docx": extract_docx,
    "txt": extract_plain,
    "md": extract_plain,
    "csv": extract_csv,
    "json": extract_json,
}


def extract_file(data: bytes, filename: str, ext: str) -> Extraction:
    fn = EXTRACTORS.get(ext)
    if fn is None:
        raise ExtractionError(f"Unsupported file type: .{ext}")
    return fn(data, filename)


def extract_text_content(content: str, name: str) -> Extraction:
    ex = Extraction()
    ex.chunks = _chunk(content, name, {"name": name})
    if not ex.chunks:
        raise ExtractionError("The text is empty")
    return ex
